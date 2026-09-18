"""Unit tests: PPTX export (slide count, titles, notes, file validity)."""
import os
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

from src.assemble import pptx_builder as PPTX  # noqa: E402


def _slides(n=2):
    out = []
    for i in range(n):
        img = np.full((90, 160, 3), 200 - i * 40, np.uint8)
        out.append({"rectified": img, "text": f"hello world slide {i}",
                    "words": [], "timestamp": float(i * 4),
                    "timestamp_start": float(i * 4)})
    return out


def test_pptx_builds_valid_file(tmp_path):
    p = str(tmp_path / "slides.pptx")
    PPTX.build_pptx(_slides(2), p)
    assert os.path.isfile(p) and os.path.getsize(p) > 1000
    from pptx import Presentation
    prs = Presentation(p)
    assert len(prs.slides) == 2


def test_pptx_titles_and_notes(tmp_path):
    p = str(tmp_path / "slides.pptx")
    PPTX.build_pptx(_slides(3), p)
    from pptx import Presentation
    prs = Presentation(p)
    for i, slide in enumerate(prs.slides):
        texts = [sh.text for sh in slide.shapes if sh.has_text_frame]
        assert any(f"Slide {i + 1}" in t for t in texts), texts
        notes = slide.notes_slide.placeholders[1].text
        assert f"slide {i}" in notes, notes
